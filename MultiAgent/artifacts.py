"""Safe, conversation-scoped Markdown, PDF, and DOCX artifact generation."""

from __future__ import annotations

import hashlib
import os
import re
import json
import zipfile
from dataclasses import dataclass
from datetime import datetime, timedelta, timezone
from pathlib import Path
from typing import Any
from xml.sax.saxutils import escape

from docx import Document
from docx.shared import Inches, Pt
from reportlab.lib import colors
from reportlab.lib.pagesizes import letter
from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
from reportlab.lib.units import inch
from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer, Preformatted, ListFlowable, ListItem

from config import PROJECT_DIR


@dataclass(frozen=True)
class GeneratedArtifact:
    name: str
    path: Path
    mime: str


def _artifact_root() -> Path:
    configured = os.getenv("APP_ARTIFACT_DIR", "data/generated")
    path = Path(configured).expanduser()
    if not path.is_absolute():
        path = PROJECT_DIR / path
    resolved = path.resolve()
    resolved.mkdir(parents=True, exist_ok=True)
    return resolved


def _slug(value: str) -> str:
    value = re.sub(r"[^A-Za-z0-9_-]+", "-", value).strip("-_").lower()
    return (value[:48] or "document")


def _conversation_folder(conversation_id: str) -> str:
    # Hash the opaque ID rather than using it as a filesystem path.
    return hashlib.sha256(conversation_id.encode("utf-8")).hexdigest()[:20]


def _cleanup(root: Path, keep_days: int = 7) -> None:
    cutoff = datetime.now(timezone.utc).timestamp() - timedelta(days=keep_days).total_seconds()
    for item in root.glob("*/*"):
        try:
            if item.is_file() and item.stat().st_mtime < cutoff:
                item.unlink()
        except OSError:
            continue


def _markdown_to_docx(content: str, path: Path) -> None:
    doc = Document()
    section = doc.sections[0]
    section.top_margin = Inches(0.65)
    section.bottom_margin = Inches(0.65)
    section.left_margin = Inches(0.75)
    section.right_margin = Inches(0.75)
    normal = doc.styles["Normal"]
    normal.font.name = "Aptos"
    normal.font.size = Pt(10)
    in_code = False
    for raw in content.splitlines():
        line = raw.rstrip()
        if line.strip().startswith("```"):
            in_code = not in_code
            continue
        if in_code:
            para = doc.add_paragraph()
            run = para.add_run(line)
            run.font.name = "Consolas"
            run.font.size = Pt(9)
        elif not line.strip():
            continue
        elif line.startswith("### "):
            doc.add_heading(line[4:], level=3)
        elif line.startswith("## "):
            doc.add_heading(line[3:], level=2)
        elif line.startswith("# "):
            doc.add_heading(line[2:], level=1)
        elif re.match(r"\s*[-*+]\s+", line):
            doc.add_paragraph(re.sub(r"^\s*[-*+]\s+", "", line), style="List Bullet")
        elif re.match(r"\s*\d+[.)]\s+", line):
            doc.add_paragraph(re.sub(r"^\s*\d+[.)]\s+", "", line), style="List Number")
        else:
            doc.add_paragraph(line.replace("**", "").replace("__", ""))
    doc.save(path)


def _markdown_to_pdf(content: str, path: Path) -> None:
    styles = getSampleStyleSheet()
    body = ParagraphStyle("DocBody", parent=styles["BodyText"], fontName="Helvetica", fontSize=9.5, leading=13, spaceAfter=6)
    code = ParagraphStyle("DocCode", parent=body, fontName="Courier", fontSize=8, leftIndent=10, backColor=colors.whitesmoke)
    story = []
    in_code = False
    code_lines: list[str] = []
    for raw in content.splitlines():
        line = raw.rstrip()
        if line.strip().startswith("```"):
            if in_code:
                story.append(Preformatted("\n".join(code_lines), code))
                story.append(Spacer(1, 6))
                code_lines = []
            in_code = not in_code
            continue
        if in_code:
            code_lines.append(line)
            continue
        if not line.strip():
            story.append(Spacer(1, 5))
            continue
        heading = re.match(r"^(#{1,3})\s+(.*)$", line)
        if heading:
            level = len(heading.group(1))
            story.append(Paragraph(escape(heading.group(2)), styles[f"Heading{level}"]))
        elif re.match(r"\s*[-*+]\s+", line):
            text = escape(re.sub(r"^\s*[-*+]\s+", "", line))
            story.append(ListFlowable([ListItem(Paragraph(text, body))], bulletType="bullet", leftIndent=16))
        else:
            clean = re.sub(r"\*\*(.*?)\*\*", r"<b>\1</b>", escape(line))
            clean = re.sub(r"\*(.*?)\*", r"<i>\1</i>", clean)
            story.append(Paragraph(clean, body))
    if in_code and code_lines:
        story.append(Preformatted("\n".join(code_lines), code))
    if not story:
        raise ValueError("Document content is empty.")
    SimpleDocTemplate(str(path), pagesize=letter, rightMargin=0.7 * inch, leftMargin=0.7 * inch, topMargin=0.65 * inch, bottomMargin=0.65 * inch, title="Generated document").build(story)


def generate_artifacts(content: str, conversation_id: str, title: str = "document", formats: tuple[str, ...] = ("pdf", "docx")) -> list[GeneratedArtifact]:
    """Generate real files, scoped to an opaque conversation hash."""
    content = content.strip()
    if not content:
        raise ValueError("Cannot generate a document from empty content.")
    normalized = tuple(dict.fromkeys(ext.lower().lstrip(".") for ext in formats))
    unsupported = set(normalized) - {"pdf", "docx", "md", "txt", "html"}
    if unsupported:
        raise ValueError("Unsupported output format.")
    root = _artifact_root()
    _cleanup(root)
    directory = root / _conversation_folder(conversation_id)
    directory.mkdir(parents=True, exist_ok=True)
    stem = _slug(title)
    artifacts = []
    for ext in normalized:
        path = directory / f"{stem}.{ext}"
        if ext == "pdf":
            _markdown_to_pdf(content, path)
            mime = "application/pdf"
        elif ext == "docx":
            _markdown_to_docx(content, path)
            mime = "application/vnd.openxmlformats-officedocument.wordprocessingml.document"
        elif ext == "md":
            path.write_text(content + "\n", encoding="utf-8")
            mime = "text/markdown"
        elif ext == "html":
            path.write_text(content, encoding="utf-8")
            mime = "text/html"
        else:
            path.write_text(content + "\n", encoding="utf-8")
            mime = "text/plain"
        if not path.is_file() or path.stat().st_size == 0:
            raise OSError("Generated file was empty.")
        artifacts.append(GeneratedArtifact(path.name, path, mime))
    return artifacts


def _safe_project_path(name: str) -> str:
    """Normalize a project-relative path and reject traversal/absolute paths."""
    value = name.replace("\\", "/").strip()
    parts = value.split("/")
    if not value or len(value) > 240 or value.startswith("/") or any(part in {"", ".", ".."} or len(part) > 100 for part in parts):
        raise ValueError("Project file path is not safe.")
    if any(not re.fullmatch(r"[A-Za-z0-9_.-]+", part) for part in parts):
        raise ValueError("Project file name contains unsupported characters.")
    return "/".join(parts)


def _contained_path(path: Path, root: Path) -> Path:
    """Resolve a path and reject symlinks that escape its trusted root."""
    resolved_root = root.resolve()
    resolved_path = path.resolve()
    if not resolved_path.is_relative_to(resolved_root):
        raise ValueError("Project path escapes its storage directory.")
    return resolved_path


def save_project(files: dict[str, str], conversation_id: str, title: str = "project") -> list[GeneratedArtifact]:
    """Persist verified source files and a ZIP in a conversation-scoped folder."""
    if not files or len(files) > 40:
        raise ValueError("Project must contain between 1 and 40 files.")
    root = _artifact_root()
    _cleanup(root)
    directory = root / _conversation_folder(conversation_id) / _slug(title)
    directory.mkdir(parents=True, exist_ok=True)
    project_root = _contained_path(directory, root)
    manifest = directory / "project.json"
    previous: dict[str, Any] = {}
    if manifest.is_file():
        try:
            _contained_path(manifest, project_root)
            loaded = json.loads(manifest.read_text(encoding="utf-8"))
            if isinstance(loaded, dict):
                previous = loaded
        except (OSError, ValueError, UnicodeError):
            previous = {}
    normalized: dict[str, str] = {}
    for name, content in files.items():
        safe_name = _safe_project_path(name)
        if safe_name.lower() == "project.json" or safe_name.lower() == f"{_slug(title)}.zip".lower():
            raise ValueError("Project file name is reserved for workspace metadata or archive output.")
        if not isinstance(content, str) or not content.strip() or len(content.encode("utf-8")) > 1_000_000:
            raise ValueError("Project files must contain nonempty text up to 1 MB each.")
        normalized[safe_name] = content
    for name, content in normalized.items():
        parent = project_root
        parts = name.split("/")
        for component in parts[:-1]:
            candidate = parent / component
            candidate.mkdir(exist_ok=True)
            parent = _contained_path(candidate, project_root)
        path = parent / parts[-1]
        _contained_path(path, project_root).write_text(content, encoding="utf-8")
        if not path.is_file() or path.read_text(encoding="utf-8") != content:
            raise OSError("Project file verification failed.")
    prior_version = max(0, int(previous.get("version", 0)))
    unchanged = previous.get("files") == normalized and prior_version > 0
    version = prior_version if unchanged else prior_version + 1
    versions = list(previous.get("versions", [])) if isinstance(previous.get("versions"), list) else []
    if not unchanged:
        version_dir = directory / "versions" / f"v{version:04d}"
        version_dir.mkdir(parents=True, exist_ok=True)
        safe_version_dir = _contained_path(version_dir, project_root)
        for name, content in normalized.items():
            destination = safe_version_dir.joinpath(*name.split("/"))
            destination.parent.mkdir(parents=True, exist_ok=True)
            _contained_path(destination, safe_version_dir).write_text(content, encoding="utf-8")
        versions.append({"version": version, "created_at": datetime.now(timezone.utc).isoformat(timespec="seconds")})
    _contained_path(manifest, project_root).write_text(
        json.dumps({"title": title, "files": normalized, "version": version, "versions": versions}, ensure_ascii=False),
        encoding="utf-8",
    )
    zip_path = directory / f"{_slug(title)}.zip"
    safe_zip_path = _contained_path(zip_path, project_root)
    with zipfile.ZipFile(safe_zip_path, "w", zipfile.ZIP_DEFLATED) as archive:
        for name, content in normalized.items():
            archive.writestr(name, content)
    if not zip_path.is_file() or zip_path.stat().st_size == 0:
        raise OSError("Project archive verification failed.")
    return [GeneratedArtifact(name, directory.joinpath(*name.split("/")), _mime_for_source(name)) for name in normalized] + [
        GeneratedArtifact(zip_path.name, zip_path, "application/zip")
    ]


def load_project(conversation_id: str) -> dict | None:
    """Load the latest project manifest for a conversation, if one exists."""
    root = _artifact_root()
    parent = root / _conversation_folder(conversation_id)
    if not parent.exists():
        return None
    try:
        _contained_path(parent, root)
    except ValueError:
        return None
    manifests = sorted(parent.glob("*/project.json"), key=lambda path: path.stat().st_mtime, reverse=True) if parent.exists() else []
    if not manifests:
        return None
    try:
        _contained_path(manifests[0], parent)
        data = json.loads(manifests[0].read_text(encoding="utf-8"))
    except (OSError, ValueError, UnicodeError):
        return None
    if not isinstance(data, dict) or not isinstance(data.get("files"), dict) or not data["files"] or len(data["files"]) > 40:
        return None
    try:
        for name, content in data["files"].items():
            _safe_project_path(name)
            if not isinstance(content, str) or len(content.encode("utf-8")) > 1_000_000:
                return None
    except (ValueError, UnicodeError):
        return None
    return {
        "title": str(data.get("title") or "project"),
        "files": data["files"],
        "version": int(data.get("version", 1)),
        "versions": data.get("versions", [{"version": 1}]),
    }


def load_project_version(conversation_id: str, version: int) -> dict | None:
    """Load a validated immutable project version for the preview history control."""
    current = load_project(conversation_id)
    if not current or version < 1 or version not in {item.get("version") for item in current["versions"] if isinstance(item, dict)}:
        return None
    root = _artifact_root()
    parent = root / _conversation_folder(conversation_id) / _slug(current["title"]) / "versions" / f"v{version:04d}"
    try:
        _contained_path(parent, root)
        files = {}
        for path in parent.rglob("*"):
            if not path.is_file():
                continue
            _contained_path(path, parent)
            name = path.relative_to(parent).as_posix()
            _safe_project_path(name)
            content = path.read_text(encoding="utf-8")
            if len(content.encode("utf-8")) > 1_000_000:
                return None
            files[name] = content
        if not files:
            return None
        return {**current, "files": files, "version": version}
    except (OSError, ValueError, UnicodeError):
        return None


def _mime_for_source(name: str) -> str:
    return {".html": "text/html", ".css": "text/css", ".js": "text/javascript", ".jsx": "text/jsx", ".tsx": "text/tsx", ".ts": "text/typescript", ".py": "text/x-python", ".json": "application/json", ".md": "text/markdown"}.get(Path(name).suffix.lower(), "text/plain")
